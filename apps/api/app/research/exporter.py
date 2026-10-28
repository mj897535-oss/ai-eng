"""Research Data Exporter — single source of truth for export schema and serialisation.

DESIGN RULE (P14):
  The DATA_DICTIONARY is generated from EXPORT_COLUMN_DEFINITIONS — the exact same
  list used by the row serialiser. They can NEVER drift from each other.

De-identification rules
-----------------------
Default export: research_code (pseudonymous), study_group, assessment_type, versions,
  counts, rates. No names, emails, audio keys, or free-text transcript fields.
Elevated export: researcher-admin-only, with explicit reason, double-confirm header,
  and a separate audit-log entry.

All rate calculations delegate to app.research.metrics.calculate_error_rate().
"""

from __future__ import annotations

import csv
import hashlib
import hmac
import io
from dataclasses import dataclass
from typing import Any

from app.research.metrics import ErrorCategory, calculate_error_rate

# ---------------------------------------------------------------------------
# Column Definition Schema — single source of truth
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class ColumnDef:
    """Typed descriptor for one export column."""

    column_name: str
    dtype: str         # "string" | "integer" | "float" | "boolean" | "date"
    definition: str
    allowed_values: str  # comma-separated list or "any"
    source_table: str
    calculation_rule: str  # "direct" or formula string
    version: str           # schema version that introduced this column
    is_identifier: bool    # True → excluded from default (de-identified) export


# ---------------------------------------------------------------------------
# Master column registry
# ---------------------------------------------------------------------------

EXPORT_COLUMN_DEFINITIONS: list[ColumnDef] = [
    # --- Participant pseudonym ---
    ColumnDef(
        column_name="research_code",
        dtype="string",
        definition=(
            "Pseudonymous participant identifier assigned at enrolment. "
            "Optionally replaced with HMAC-SHA256 hash when hash_participant_ids=true."
        ),
        allowed_values="any",
        source_table="research.participants",
        calculation_rule="direct",
        version="1.0",
        is_identifier=False,
    ),
    # --- Assessment metadata ---
    ColumnDef(
        column_name="assessment_id",
        dtype="string",
        definition="UUID of the assessment record.",
        allowed_values="any",
        source_table="research.assessments",
        calculation_rule="direct",
        version="1.0",
        is_identifier=False,
    ),
    ColumnDef(
        column_name="study_group",
        dtype="string",
        definition="Participant randomised study group.",
        allowed_values="intervention, control",
        source_table="research.participants",
        calculation_rule="direct",
        version="1.0",
        is_identifier=False,
    ),
    ColumnDef(
        column_name="assessment_type",
        dtype="string",
        definition="Whether this is a pre-test or post-test assessment.",
        allowed_values="pre, post",
        source_table="research.assessments",
        calculation_rule="direct",
        version="1.0",
        is_identifier=False,
    ),
    ColumnDef(
        column_name="assessment_date",
        dtype="date",
        definition=(
            "Date (not datetime) the assessment was created. "
            "Time component removed to reduce re-identification risk."
        ),
        allowed_values="any",
        source_table="research.assessments",
        calculation_rule="date(assessment.created_at)",
        version="1.0",
        is_identifier=False,
    ),
    ColumnDef(
        column_name="status",
        dtype="string",
        definition="Processing status of the assessment at export time.",
        allowed_values="completed, pending, failed, needs_attention",
        source_table="research.assessments",
        calculation_rule="direct",
        version="1.0",
        is_identifier=False,
    ),
    ColumnDef(
        column_name="is_verified",
        dtype="boolean",
        definition=(
            "True when all AI candidate annotations have been resolved by a human reviewer. "
            "Unverified assessments are marked provisional."
        ),
        allowed_values="true, false",
        source_table="research.assessments",
        calculation_rule="direct",
        version="1.0",
        is_identifier=False,
    ),
    ColumnDef(
        column_name="is_official",
        dtype="boolean",
        definition="True when the assessment is designated official for research purposes.",
        allowed_values="true, false",
        source_table="research.assessments",
        calculation_rule="direct",
        version="1.0",
        is_identifier=False,
    ),
    # --- Word / timing counts ---
    ColumnDef(
        column_name="word_count",
        dtype="integer",
        definition=(
            "Official word count computed by the frozen word-counting rule "
            "recorded in rubric_version."
        ),
        allowed_values="any",
        source_table="research.assessments",
        calculation_rule="direct (stored at transcription time)",
        version="1.0",
        is_identifier=False,
    ),
    ColumnDef(
        column_name="audio_duration_seconds",
        dtype="float",
        definition="Duration of the audio recording in seconds.",
        allowed_values="any",
        source_table="research.assessments",
        calculation_rule="direct",
        version="1.0",
        is_identifier=False,
    ),
    # --- Version provenance ---
    ColumnDef(
        column_name="rubric_version",
        dtype="string",
        definition="Version tag of the rubric used for this assessment.",
        allowed_values="any",
        source_table="research.assessments",
        calculation_rule="direct",
        version="1.0",
        is_identifier=False,
    ),
    ColumnDef(
        column_name="stt_version",
        dtype="string",
        definition="Version tag of the speech-to-text provider/model used.",
        allowed_values="any",
        source_table="research.assessments",
        calculation_rule="direct",
        version="1.0",
        is_identifier=False,
    ),
    ColumnDef(
        column_name="prompt_version",
        dtype="string",
        definition="Version tag of the LLM assessor prompt used.",
        allowed_values="any",
        source_table="research.assessments",
        calculation_rule="direct",
        version="1.0",
        is_identifier=False,
    ),
    ColumnDef(
        column_name="model_version",
        dtype="string",
        definition="Version tag of the LLM model used for error analysis.",
        allowed_values="any",
        source_table="research.assessments",
        calculation_rule="direct",
        version="1.0",
        is_identifier=False,
    ),
    # --- Per-category error counts (effective annotation set) ---
    *[
        ColumnDef(
            column_name=f"count_{cat.value.lower()}",
            dtype="integer",
            definition=(
                f"Number of assessable {cat.value} errors in the effective annotation set "
                f"(reviewer decisions override AI; uncertain/transcription_related/"
                f"acceptable_variation excluded)."
            ),
            allowed_values="any",
            source_table="research.error_annotations (derived)",
            calculation_rule=(
                "COUNT(effective_annotations WHERE category_code='"
                + cat.value
                + "' AND assessable=true)"
            ),
            version="1.0",
            is_identifier=False,
        )
        for cat in ErrorCategory
    ],
    # --- Total assessable error count ---
    ColumnDef(
        column_name="total_assessable_errors",
        dtype="integer",
        definition="Sum of all assessable errors across all categories.",
        allowed_values="any",
        source_table="research.error_annotations (derived)",
        calculation_rule="SUM(count_tns + count_sva + count_wo + count_aux + count_art + count_prep + count_pro + count_oth)",
        version="1.0",
        is_identifier=False,
    ),
    # --- Overall error rate ---
    ColumnDef(
        column_name="error_rate_per_100_words",
        dtype="float",
        definition="Overall error rate: (total_assessable_errors / word_count) * 100.",
        allowed_values="any",
        source_table="derived",
        calculation_rule="calculate_error_rate(total_assessable_errors, word_count) — app.research.metrics",
        version="1.0",
        is_identifier=False,
    ),
    # --- Per-category rates ---
    *[
        ColumnDef(
            column_name=f"rate_{cat.value.lower()}_per_100",
            dtype="float",
            definition=(
                f"Error rate for {cat.value} per 100 words: "
                f"(count_{cat.value.lower()} / word_count) * 100."
            ),
            allowed_values="any",
            source_table="derived",
            calculation_rule=f"calculate_error_rate(count_{cat.value.lower()}, word_count) — app.research.metrics",
            version="1.0",
            is_identifier=False,
        )
        for cat in ErrorCategory
    ],
    # --- Provisional flag ---
    ColumnDef(
        column_name="provisional",
        dtype="boolean",
        definition="True when the assessment has not been fully verified by a human reviewer.",
        allowed_values="true, false",
        source_table="derived",
        calculation_rule="NOT is_verified",
        version="1.0",
        is_identifier=False,
    ),
]

# Derived sub-lists used by the exporter
SAFE_COLUMNS: list[ColumnDef] = [c for c in EXPORT_COLUMN_DEFINITIONS if not c.is_identifier]
ELEVATED_COLUMNS: list[ColumnDef] = EXPORT_COLUMN_DEFINITIONS  # all columns

# Safe column names (for header-list assertions in tests)
SAFE_COLUMN_NAMES: list[str] = [c.column_name for c in SAFE_COLUMNS]
ELEVATED_COLUMN_NAMES: list[str] = [c.column_name for c in ELEVATED_COLUMNS]

# Columns that must NEVER appear in a default export
FORBIDDEN_DEFAULT_COLUMNS: set[str] = {
    "full_name", "email", "audio_key", "transcript", "raw_stt_response_key",
    "user_id", "mother_tongue", "school_medium", "country",
}


# ---------------------------------------------------------------------------
# Data Dictionary generators
# ---------------------------------------------------------------------------

def generate_data_dictionary_csv(columns: list[ColumnDef]) -> bytes:
    """Generate DATA_DICTIONARY CSV bytes from the given column list."""
    buf = io.StringIO()
    writer = csv.DictWriter(
        buf,
        fieldnames=[
            "column_name", "dtype", "definition", "allowed_values",
            "source_table", "calculation_rule", "version",
        ],
        lineterminator="\n",
    )
    writer.writeheader()
    for col in columns:
        writer.writerow({
            "column_name": col.column_name,
            "dtype": col.dtype,
            "definition": col.definition,
            "allowed_values": col.allowed_values,
            "source_table": col.source_table,
            "calculation_rule": col.calculation_rule,
            "version": col.version,
        })
    return buf.getvalue().encode("utf-8")


def generate_data_dictionary_md(columns: list[ColumnDef]) -> str:
    """Generate DATA_DICTIONARY Markdown table from the given column list."""
    header = (
        "| column_name | dtype | definition | allowed_values "
        "| source_table | calculation_rule | version |\n"
        "|---|---|---|---|---|---|---|\n"
    )
    rows = []
    for col in columns:
        rows.append(
            f"| `{col.column_name}` | {col.dtype} "
            f"| {col.definition.replace('|', '/')} "
            f"| {col.allowed_values} "
            f"| {col.source_table} "
            f"| {col.calculation_rule} "
            f"| {col.version} |"
        )
    return header + "\n".join(rows)


# ---------------------------------------------------------------------------
# Participant ID hashing
# ---------------------------------------------------------------------------

def hash_participant_id(research_code: str, salt: str) -> str:
    """HMAC-SHA256 of research_code with salt.  Same code → same hash within a snapshot."""
    return hmac.new(
        salt.encode("utf-8"),
        research_code.encode("utf-8"),
        hashlib.sha256,
    ).hexdigest()

# ---------------------------------------------------------------------------
# k-Anonymity check
# ---------------------------------------------------------------------------

def compute_k_anonymity_warnings(
    rows: list[dict[str, Any]],
    quasi_identifiers: list[str],
    k_threshold: int = 5,
) -> list[str]:
    """Return warning strings for quasi-identifier combinations with n < k_threshold."""
    from collections import Counter
    combo_counts: Counter = Counter()
    for row in rows:
        key = tuple(str(row.get(qi, "")) for qi in quasi_identifiers)
        combo_counts[key] += 1

    warnings = []
    for combo, count in combo_counts.items():
        if count < k_threshold:
            label = ", ".join(f"{qi}={val}" for qi, val in zip(quasi_identifiers, combo, strict=False))
            warnings.append(
                f"k-anonymity warning: subgroup ({label}) has only {count} record(s) "
                f"(threshold={k_threshold}). Consider suppressing or generalising before sharing."
            )
    return warnings


# ---------------------------------------------------------------------------
# Row builder
# ---------------------------------------------------------------------------

_CAT_CODES = [cat.value for cat in ErrorCategory]


def build_export_rows(
    assessments: list[dict[str, Any]],
    category_counts_by_assessment: dict[str, dict[str, int]],
    include_identifiers: bool = False,
    salt: str | None = None,
) -> list[dict[str, Any]]:
    """Build export rows from assessment records.

    Parameters
    ----------
    assessments:
        List of dicts, each containing fields from Assessment + Participant.
    category_counts_by_assessment:
        Dict mapping assessment_id (str) to {category_code: count} from effective annotation set.
    include_identifiers:
        When True (elevated export), include identifier-flagged columns.
    salt:
        When provided, HMAC-hash the research_code field using this salt.
    """
    columns = ELEVATED_COLUMNS if include_identifiers else SAFE_COLUMNS
    col_names = {c.column_name for c in columns}
    rows: list[dict[str, Any]] = []

    for asmt in assessments:
        aid = str(asmt.get("assessment_id", ""))
        cat_counts = category_counts_by_assessment.get(aid, {})
        word_count = int(asmt.get("word_count") or 0)

        # Per-category counts
        per_cat_counts: dict[str, int] = {cat: cat_counts.get(cat, 0) for cat in _CAT_CODES}
        total_errors = sum(per_cat_counts.values())

        # research_code — optionally hashed
        rc = asmt.get("research_code", "")
        if salt:
            rc = hash_participant_id(rc, salt)

        row: dict[str, Any] = {
            "research_code": rc,
            "assessment_id": aid,
            "study_group": asmt.get("study_group", ""),
            "assessment_type": asmt.get("assessment_type", ""),
            "assessment_date": str(asmt.get("assessment_date", ""))[:10],
            "status": asmt.get("status", ""),
            "is_verified": asmt.get("is_verified", False),
            "is_official": asmt.get("is_official", False),
            "word_count": word_count,
            "audio_duration_seconds": asmt.get("audio_duration_seconds", 0.0),
            "rubric_version": asmt.get("rubric_version", ""),
            "stt_version": asmt.get("stt_version", ""),
            "prompt_version": asmt.get("prompt_version", ""),
            "model_version": asmt.get("model_version", ""),
            # Per-category counts
            **{f"count_{cat.lower()}": per_cat_counts[cat] for cat in _CAT_CODES},
            "total_assessable_errors": total_errors,
            "error_rate_per_100_words": calculate_error_rate(total_errors, word_count),
            # Per-category rates — use metrics module (never re-implement)
            **{
                f"rate_{cat.lower()}_per_100": calculate_error_rate(per_cat_counts[cat], word_count)
                for cat in _CAT_CODES
            },
            "provisional": not asmt.get("is_verified", False),
        }

        # Filter to allowed columns only
        rows.append({k: v for k, v in row.items() if k in col_names})

    return rows


# ---------------------------------------------------------------------------
# CSV / XLSX serialisers
# ---------------------------------------------------------------------------

def generate_csv_bytes(rows: list[dict[str, Any]], columns: list[ColumnDef]) -> bytes:
    """Serialise rows to CSV bytes. Uses only column order from columns list."""
    if not columns:
        return b""
    fieldnames = [c.column_name for c in columns]
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=fieldnames, lineterminator="\n", extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def generate_xlsx_bytes(rows: list[dict[str, Any]], columns: list[ColumnDef]) -> bytes:
    """Serialise rows to XLSX bytes using openpyxl."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for XLSX export. Install it first.") from exc

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SyntaxLens Export"

    fieldnames = [c.column_name for c in columns]
    header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx, name in enumerate(fieldnames, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, name in enumerate(fieldnames, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(name, ""))

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Checksum
# ---------------------------------------------------------------------------

def compute_sha256(data: bytes) -> str:
    """Return hex SHA-256 digest of data."""
    return hashlib.sha256(data).hexdigest()
